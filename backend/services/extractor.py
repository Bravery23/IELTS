"""
Text extraction service for PDF and DOCX files.
Extracts raw text content, ignoring images and diagrams.
"""
import io
from fastapi import UploadFile
import pdfplumber
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file using pdfplumber."""
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            # Preserve basic formatting info
            runs_text = []
            for run in paragraph.runs:
                text = run.text
                if run.bold and text.strip():
                    text = f"<b>{text}</b>"
                if run.italic and text.strip():
                    text = f"<i>{text}</i>"
                runs_text.append(text)
            if runs_text:
                text_parts.append("".join(runs_text))
            else:
                text_parts.append(paragraph.text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n\n".join(text_parts)


async def extract_text(file: UploadFile) -> str:
    """
    Route file to the appropriate extractor based on file extension.
    Supports .pdf and .docx files.
    """
    filename = file.filename or ""
    file_bytes = await file.read()

    if filename.lower().endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.lower().endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file format: {filename}. Please upload a .pdf or .docx file."
        )
